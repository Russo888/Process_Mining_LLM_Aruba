"""Convert the maintained project report to an editable Word document."""
import re
import argparse
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser()
parser.add_argument('--source', default='DOCUMENTAZIONE_PROGETTO.md')
parser.add_argument('--output', default='DOCUMENTAZIONE_PROGETTO.docx')
parser.add_argument('--language', choices=['it', 'en'], default='it')
args = parser.parse_args()
source_text = (ROOT/args.source).read_text(encoding='utf-8')
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(2)
sec.left_margin = sec.right_margin = Cm(1.8)
sec.footer_distance = Cm(1)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet', 'List Number']:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0, 0, 0)
    language = OxmlElement('w:lang')
    language.set(qn('w:val'), 'en-GB' if args.language == 'en' else 'it-IT')
    st.element.get_or_add_rPr().append(language)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(25)
for name, size in [('Heading 1', 16), ('Heading 2', 12.5)]:
    st = doc.styles[name]
    st.font.size = Pt(size)
    st.paragraph_format.space_before = Pt(15)
    st.paragraph_format.space_after = Pt(7)
    st.paragraph_format.keep_with_next = True
doc.styles['Caption'].font.name = 'Calibri'
doc.styles['Caption'].font.size = Pt(9)
doc.styles['Caption'].font.color.rgb = RGBColor(0, 0, 0)

# The default Word template can add a blue rule beneath Title through its style.
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)

def inline(p, text):
    # Preserve evidence references as portable repository paths in the document.
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', text)
    for part in re.split(r'(\*\*.*?\*\*|`[^`]+`)', text):
        if not part:
            continue
        r = p.add_run(part[2:-2] if part.startswith('**') else part[1:-1] if part.startswith('`') else part)
        if part.startswith('**'):
            r.bold = True
        if part.startswith('`'):
            r.font.name = 'Consolas'
            r.font.size = Pt(9)

def table(lines):
    values = [[c.strip() for c in line.strip().strip('|').split('|')] for line in lines]
    values = [v for v in values if not all(re.fullmatch(r':?-+:?', x) for x in v)]
    cols = len(values[0])
    t = doc.add_table(rows=0, cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    ratios = {2:[.27,.73], 3:[.40,.30,.30], 4:[.24,.24,.24,.28], 7:[.22,.16,.11,.12,.14,.14,.11]}[cols]
    if values[0][0] in ('Percorso nel progetto', 'Project location'):
        ratios = [.45, .55]
    for col, ratio in zip(t.columns, ratios):
        col.width = Cm(17.4*ratio)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge)
        for key,value in [('val','single'),('sz','4'),('color','D9D9D9')]: e.set(qn('w:'+key),value)
        borders.append(e)
    pr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side in ['top','bottom','left','right']:
        e=OxmlElement('w:'+side); e.set(qn('w:w'),'85'); e.set(qn('w:type'),'dxa'); margins.append(e)
    pr.append(margins)
    for index, cells in enumerate(values):
        row=t.add_row()
        no_split=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
        if index==0:
            header=OxmlElement('w:tblHeader'); row._tr.get_or_add_trPr().append(header)
        for j, (cell, value) in enumerate(zip(row.cells,cells)):
            cell.width=Cm(17.4*ratios[j]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'DCE6EF' if index==0 else ('F5F7F9' if index%2==0 else 'FFFFFF'))
            cell._tc.get_or_add_tcPr().append(shade)
            p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(2); p.paragraph_format.space_before=Pt(2)
            if index == 0:
                p.paragraph_format.keep_with_next = True
            p.paragraph_format.line_spacing=1.0
            inline(p,value)
            for r in p.runs:
                r.font.size=Pt(9 if cols==7 else 9.5)
                if index==0: r.bold=True
            if j>0 and (cols==7 or re.fullmatch(r'[\d.,/% —]+',value)):
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

lines=source_text.splitlines()
i=0
list_number_id = None
while i<len(lines):
    line=lines[i]
    if not line.strip(): i+=1; continue
    if line.startswith('|'):
        block=[]
        while i<len(lines) and lines[i].startswith('|'): block.append(lines[i]); i+=1
        table(block); continue
    image = re.fullmatch(r'!\[([^\]]+)\]\(([^)]+)\)', line)
    if image:
        image_path = ROOT/image.group(2)
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        picture = p.add_run().add_picture(str(image_path), width=Cm(17.2))
        picture._inline.docPr.set('descr', image.group(1))
        doc.add_paragraph(image.group(1), 'Caption')
    elif re.fullmatch(r'\[SCREENSHOT S[1-8]\]', line):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(36)
        p.paragraph_format.keep_with_next = True
        p.add_run(line).bold = True
    elif line.startswith('#'):
        marker, content=line.split(' ',1)
        if len(marker) > 2:
            content = re.sub(r'^\d+(?:\.\d+)*\.?\s+', '', content)
        content=re.sub(r'[^\w\s]', ' ', content)
        content=re.sub(r'\s+',' ',content).strip()
        doc.add_paragraph(content,'Title' if len(marker)==1 else 'Heading '+str(len(marker)-1))
    elif line.startswith('- '): inline(doc.add_paragraph(style='List Bullet'),line[2:])
    elif args.language == 'en' and re.match(r'^\d+\. ',line):
        # Explicit labels avoid Word's context-dependent list restarts across pages.
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(.63)
        p.paragraph_format.first_line_indent = Cm(-.63)
        p.paragraph_format.space_after = Pt(3)
        inline(p, line)
    elif re.match(r'^\d+\. ',line):
        if line.startswith('1. ') or list_number_id is None:
            numbering = doc.part.numbering_part.element
            style_num_id = doc.styles['List Number'].element.pPr.numPr.numId.val
            base_num = next(n for n in numbering.findall(qn('w:num')) if n.get(qn('w:numId')) == str(style_num_id))
            abstract_id = base_num.find(qn('w:abstractNumId')).get(qn('w:val'))
            list_number_id = max(int(n.get(qn('w:numId'))) for n in numbering.findall(qn('w:num'))) + 1
            num = OxmlElement('w:num'); num.set(qn('w:numId'), str(list_number_id))
            abstract = OxmlElement('w:abstractNumId'); abstract.set(qn('w:val'), abstract_id); num.append(abstract)
            override = OxmlElement('w:lvlOverride'); override.set(qn('w:ilvl'), '0')
            restart = OxmlElement('w:startOverride'); restart.set(qn('w:val'), '1'); override.append(restart)
            num.append(override); numbering.append(num)
        p = doc.add_paragraph(style='List Number')
        num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = 0
        num_pr.get_or_add_numId().val = list_number_id
        inline(p,re.sub(r'^\d+\. ','',line))
    else: inline(doc.add_paragraph(),line)
    i+=1
footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Page ' if args.language == 'en' else 'Pagina ').font.size=Pt(9)
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
doc.core_properties.title=lines[0].removeprefix('# ')
doc.core_properties.subject=('Experimental development and project results' if args.language == 'en'
                            else 'Percorso sperimentale e risultati del progetto')
doc.core_properties.author=''
doc.core_properties.keywords='Aruba, process mining, esperimenti, Colab'
out=ROOT/args.output
try:
    doc.save(out)
except PermissionError:
    if args.output != 'DOCUMENTAZIONE_PROGETTO.docx':
        raise
    out = ROOT/'DOCUMENTAZIONE_PROGETTO_CON_GUIDA.docx'
    doc.save(out)
check=Document(out)
assert len(check.tables)==sum(line.startswith('|---') for line in lines), len(check.tables)
assert len(check.inline_shapes)==sum(bool(re.fullmatch(r'!\[([^\]]+)\]\(([^)]+)\)',line)) for line in lines)
assert sum(bool(re.fullmatch(r'\[SCREENSHOT S[1-8]\]', p.text)) for p in check.paragraphs)==sum(bool(re.fullmatch(r'\[SCREENSHOT S[1-8]\]',line)) for line in lines)
print(out)
print('Paragraphs:',len(check.paragraphs),'Tables:',len(check.tables))
